import openpyxl
from openpyxl import load_workbook 
from openpyxl.chart import BarChart,Reference
objWorkbook = openpyxl.load_workbook('Student Table.xlsx')
objWorksheet = objWorkbook.worksheets[0]
values = Reference(objWorksheet,min_col=2, min_row=5, max_col=6, max_row=5)
titles = Reference(objWorksheet,min_col=2, min_row=1, max_row=1, max_col=6)
chart = BarChart()
chart.add_data(values,titles)
chart.set_categories(titles) 
objWorksheet.add_chart(chart, "H7")
objWorkbook.save("Student Table.xlsx")
